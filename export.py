import os
import time
from openpyxl import Workbook, load_workbook
from models import QualifiedLead


def append_lead_to_excel(lead: QualifiedLead, file_path: str):
    headers = [
        "数据来源",
        "平台展示名称",
        "法定注册公司",
        "公司注册地址",
        "联系人",
        "联系人职位",
        "联系电话",
        "电子邮箱",
        "独立官网",
        "主营品类",
        "ICP评分",
        "核心优势总结"
    ]

    target_path = file_path
    if not os.path.exists(target_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "客户线索池"
        ws.append(headers)
    else:
        try:
            wb = load_workbook(target_path)
            ws = wb.active
        except Exception:
            wb = Workbook()
            ws = wb.active
            ws.title = "客户线索池"
            ws.append(headers)

    ws.append([
        lead.data_source or "Global Sources",
        lead.company_name,
        lead.registered_company or "未找到",
        lead.registered_address or "未找到",
        lead.contact_person or "未找到",
        lead.contact_title or "未找到",
        lead.phone or "未找到",
        lead.email or "未找到",
        lead.official_website or "未找到",
        lead.main_products or "未提及",
        lead.icp_score,
        lead.core_competence or "无"
    ])

    try:
        wb.save(target_path)
    except PermissionError:
        backup_path = f"leads_backup_{int(time.time())}.xlsx"
        wb.save(backup_path)
        print(f"⚠️ 文件被占用，已保存备用副本: {backup_path}")